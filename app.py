import io
import re
from typing import List, Tuple

import pandas as pd
import streamlit as st
from groq import Groq
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# ── Config ──────────────────────────────────────────────────────────────────
import os

GROQ_API_KEY = st.secrets["GROQ_API_KEY"]
MODEL_ID = "llama-3.3-70b-versatile"
CHUNK_SIZE = 900
CHUNK_OVERLAP = 150
TOP_K_CHUNKS = 5

# ── Page setup ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="DocuMind AI",
    page_icon="📚",
    layout="wide",
    initial_sidebar_state="expanded",
)

CUSTOM_CSS = """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

    *, *::before, *::after {
        box-sizing: border-box;
    }

    html, body, [class*="css"] {
        font-family: 'Inter', sans-serif;
    }

    /* ── Animated Header ── */
    .main-header {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 50%, #f093fb 100%);
        background-size: 200% 200%;
        animation: gradientShift 6s ease infinite;
        padding: 2.25rem 2.5rem;
        border-radius: 20px;
        margin-bottom: 1.5rem;
        color: white;
        position: relative;
        overflow: hidden;
        box-shadow:
            0 20px 60px rgba(102, 126, 234, 0.4),
            0 4px 16px rgba(102, 126, 234, 0.2);
    }

    /* Floating orb decorations */
    .main-header::before {
        content: '';
        position: absolute;
        top: -40px;
        right: -40px;
        width: 180px;
        height: 180px;
        border-radius: 50%;
        background: rgba(255,255,255,0.08);
        animation: floatOrb 8s ease-in-out infinite;
    }

    .main-header::after {
        content: '';
        position: absolute;
        bottom: -60px;
        right: 120px;
        width: 120px;
        height: 120px;
        border-radius: 50%;
        background: rgba(255,255,255,0.06);
        animation: floatOrb 6s ease-in-out infinite reverse;
    }

    .main-header h1 {
        margin: 0;
        font-size: 2rem;
        font-weight: 800;
        letter-spacing: -0.75px;
        position: relative;
        z-index: 1;
        text-shadow: 0 2px 12px rgba(0,0,0,0.15);
    }

    .main-header p {
        margin: 0.5rem 0 0 0;
        opacity: 0.88;
        font-size: 1rem;
        position: relative;
        z-index: 1;
    }

    /* ── Stat Cards ── */
    .stat-card {
        background: #ffffff;
        border: 1px solid rgba(102, 126, 234, 0.15);
        border-radius: 16px;
        padding: 1.25rem 1.5rem;
        text-align: center;
        position: relative;
        overflow: hidden;
        transition: transform 0.25s ease, box-shadow 0.25s ease, border-color 0.25s ease;
        cursor: default;
    }

    /* Shimmer sweep on hover */
    .stat-card::before {
        content: '';
        position: absolute;
        inset: 0;
        background: linear-gradient(
            105deg,
            transparent 40%,
            rgba(102,126,234,0.08) 50%,
            transparent 60%
        );
        transform: translateX(-100%);
        transition: transform 0s;
    }

    .stat-card:hover::before {
        transform: translateX(100%);
        transition: transform 0.5s ease;
    }

    .stat-card:hover {
        transform: translateY(-4px);
        box-shadow: 0 12px 32px rgba(102,126,234,0.2);
        border-color: rgba(102,126,234,0.35);
    }

    .stat-card .value {
        font-size: 1.75rem;
        font-weight: 700;
        background: linear-gradient(135deg, #667eea, #764ba2);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        background-clip: text;
        display: block;
        animation: countUp 0.6s ease both;
    }

    .stat-card .label {
        font-size: 0.78rem;
        color: #6b7280;
        margin-top: 0.3rem;
        text-transform: uppercase;
        letter-spacing: 0.5px;
        font-weight: 500;
    }

    /* ── Doc Pills ── */
    .doc-pill {
        display: inline-flex;
        align-items: center;
        gap: 0.3rem;
        background: linear-gradient(135deg, #ede9fe, #e0e7ff);
        color: #5b21b6;
        padding: 0.3rem 0.85rem;
        border-radius: 999px;
        font-size: 0.8rem;
        margin: 0.2rem;
        font-weight: 600;
        border: 1px solid rgba(91,33,182,0.15);
        transition: transform 0.2s ease, background 0.2s ease, box-shadow 0.2s ease;
        cursor: default;
    }

    .doc-pill:hover {
        transform: scale(1.06) translateY(-1px);
        background: linear-gradient(135deg, #ddd6fe, #c7d2fe);
        box-shadow: 0 4px 12px rgba(91,33,182,0.18);
    }

    /* ── Sidebar ── */
    div[data-testid="stSidebar"] {
        background: linear-gradient(180deg, #fafbff 0%, #f1f3fb 100%) !important;
        border-right: 1px solid rgba(102,126,234,0.1);
    }

    div[data-testid="stSidebar"] .stFileUploader {
        border: 2px dashed #c7d2fe;
        border-radius: 14px;
        padding: 0.75rem;
        transition: border-color 0.2s ease, background 0.2s ease;
    }

    div[data-testid="stSidebar"] .stFileUploader:hover {
        border-color: #818cf8;
        background: rgba(129,140,248,0.04);
    }

    /* ── Chat Messages ── */
    .stChatMessage {
        border-radius: 16px !important;
        animation: slideUp 0.3s ease both;
        transition: box-shadow 0.2s ease;
    }

    .stChatMessage:hover {
        box-shadow: 0 4px 20px rgba(0,0,0,0.06);
    }

    /* ── Buttons ── */
    .stButton > button {
        background: linear-gradient(135deg, #667eea, #764ba2) !important;
        background-size: 200% 200% !important;
        animation: gradientShift 4s ease infinite !important;
        color: white !important;
        border: none !important;
        border-radius: 10px !important;
        font-weight: 600 !important;
        letter-spacing: 0.3px !important;
        transition: transform 0.2s ease, box-shadow 0.2s ease !important;
    }

    .stButton > button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 8px 20px rgba(102,126,234,0.4) !important;
    }

    .stButton > button:active {
        transform: translateY(0) scale(0.97) !important;
    }

    /* ── Progress / Spinner ── */
    .stSpinner > div {
        border-top-color: #667eea !important;
    }

    /* ── Keyframes ── */
    @keyframes gradientShift {
        0%   { background-position: 0% 50%; }
        50%  { background-position: 100% 50%; }
        100% { background-position: 0% 50%; }
    }

    @keyframes floatOrb {
        0%, 100% { transform: translate(0, 0) scale(1); }
        33%       { transform: translate(-15px, 10px) scale(1.05); }
        66%       { transform: translate(10px, -15px) scale(0.95); }
    }

    @keyframes slideUp {
        from { opacity: 0; transform: translateY(10px); }
        to   { opacity: 1; transform: translateY(0); }
    }

    @keyframes countUp {
        from { opacity: 0; transform: scale(0.85); }
        to   { opacity: 1; transform: scale(1); }
    }

    /* ── Scrollbar ── */
    ::-webkit-scrollbar { width: 6px; }
    ::-webkit-scrollbar-track { background: transparent; }
    ::-webkit-scrollbar-thumb {
        background: linear-gradient(180deg, #667eea, #764ba2);
        border-radius: 99px;
    }

    footer    { visibility: hidden; }
    #MainMenu { visibility: hidden; }
</style>
"""

st.markdown(CUSTOM_CSS, unsafe_allow_html=True)


# ── Document parsers ─────────────────────────────────────────────────────────
def extract_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_doc(file_bytes: bytes) -> str:
    """Best-effort extraction for legacy .doc binary files."""
    strings = re.findall(rb"[\x20-\x7e\r\n\t]{6,}", file_bytes)
    text = b" ".join(strings).decode("ascii", errors="ignore")
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) < 50:
        raise ValueError(
            "Could not read this .doc file. Please save it as .docx and re-upload."
        )
    return text


def extract_excel(file_bytes: bytes, filename: str) -> str:
    engine = "openpyxl" if filename.lower().endswith(".xlsx") else "xlrd"
    sheets = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None, engine=engine)
    parts = []
    for name, df in sheets.items():
        df = df.fillna("")
        parts.append(f"=== Sheet: {name} ===\n{df.to_string(index=False)}")
    return "\n\n".join(parts)


def extract_csv(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            df = pd.read_csv(io.BytesIO(file_bytes), encoding=encoding)
            return df.fillna("").to_string(index=False)
        except (UnicodeDecodeError, pd.errors.ParserError):
            continue
    df = pd.read_csv(io.BytesIO(file_bytes), encoding="utf-8", errors="ignore")
    return df.fillna("").to_string(index=False)


def parse_uploaded_file(filename: str, file_bytes: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1]
    parsers = {
        "txt": extract_txt,
        "pdf": extract_pdf,
        "docx": extract_docx,
        "doc": extract_doc,
        "xlsx": lambda b: extract_excel(b, filename),
        "xls": lambda b: extract_excel(b, filename),
        "csv": extract_csv,
    }
    parser = parsers.get(ext)
    if not parser:
        raise ValueError(f"Unsupported file type: .{ext}")
    text = parser(file_bytes)
    if not text or not text.strip():
        raise ValueError(f"No readable text found in {filename}")
    return text.strip()


# ── Chunking & retrieval ─────────────────────────────────────────────────────
def chunk_text(text: str, source: str) -> List[dict]:
    chunks = []
    start = 0
    while start < len(text):
        end = start + CHUNK_SIZE
        chunk = text[start:end]
        if chunk.strip():
            chunks.append({"text": chunk.strip(), "source": source})
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks


def retrieve_relevant_chunks(query: str, chunks: List[dict], top_k: int = TOP_K_CHUNKS) -> List[dict]:
    if not chunks:
        return []
    if len(chunks) <= top_k:
        return chunks

    corpus = [c["text"] for c in chunks]
    vectorizer = TfidfVectorizer(stop_words="english", max_features=5000)
    try:
        tfidf_matrix = vectorizer.fit_transform(corpus + [query])
    except ValueError:
        return chunks[:top_k]

    query_vec = tfidf_matrix[-1]
    doc_vecs = tfidf_matrix[:-1]
    scores = cosine_similarity(query_vec, doc_vecs).flatten()
    top_indices = scores.argsort()[-top_k:][::-1]
    return [chunks[i] for i in top_indices if scores[i] > 0] or chunks[:top_k]


def build_context(chunks: List[dict]) -> str:
    if not chunks:
        return ""
    parts = []
    for i, chunk in enumerate(chunks, 1):
        parts.append(f"[Source: {chunk['source']} | Chunk {i}]\n{chunk['text']}")
    return "\n\n---\n\n".join(parts)


# ── Groq chat ────────────────────────────────────────────────────────────────
def get_groq_client() -> Groq:
    return Groq(api_key=GROQ_API_KEY)


def ask_groq(messages: List[dict]) -> str:
    client = get_groq_client()
    response = client.chat.completions.create(
        model=MODEL_ID,
        messages=messages,
        temperature=0.3,
        max_tokens=2048,
    )
    return response.choices[0].message.content or ""


# ── Session state ────────────────────────────────────────────────────────────
def init_session_state():
    defaults = {
        "messages": [],
        "chunks": [],
        "loaded_files": [],
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


init_session_state()


# ── Sidebar ──────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("### 📂 Upload Documents")
    st.caption("TXT · PDF · DOC · DOCX · Excel · CSV")

    uploaded_files = st.file_uploader(
        "Drop files here",
        type=["txt", "pdf", "doc", "docx", "xlsx", "xls", "csv"],
        accept_multiple_files=True,
        label_visibility="collapsed",
    )

    if uploaded_files:
        if st.button("🔄 Process Files", use_container_width=True, type="primary"):
            new_chunks = []
            new_names = []
            errors = []

            with st.spinner("Reading your documents..."):
                for uploaded in uploaded_files:
                    try:
                        file_bytes = uploaded.read()
                        text = parse_uploaded_file(uploaded.name, file_bytes)
                        file_chunks = chunk_text(text, uploaded.name)
                        new_chunks.extend(file_chunks)
                        new_names.append(uploaded.name)
                    except Exception as exc:
                        errors.append(f"**{uploaded.name}**: {exc}")

            if new_chunks:
                st.session_state.chunks = new_chunks
                st.session_state.loaded_files = new_names
                st.session_state.messages = []
                st.success(f"Loaded {len(new_names)} file(s) · {len(new_chunks)} chunks")
            if errors:
                for err in errors:
                    st.error(err)

    st.divider()

    if st.session_state.loaded_files:
        st.markdown("**Loaded files**")
        for name in st.session_state.loaded_files:
            st.markdown(f'<span class="doc-pill">📄 {name}</span>', unsafe_allow_html=True)

        col1, col2 = st.columns(2)
        with col1:
            st.metric("Chunks", len(st.session_state.chunks))
        with col2:
            total_chars = sum(len(c["text"]) for c in st.session_state.chunks)
            st.metric("Characters", f"{total_chars:,}")

        if st.button("🗑️ Clear All", use_container_width=True):
            st.session_state.messages = []
            st.session_state.chunks = []
            st.session_state.loaded_files = []
            st.rerun()
    else:
        st.info("Upload documents to start chatting.")

    st.divider()
    st.markdown("**💡 Tips**")
    st.markdown(
        "- Upload one or more files\n"
        "- Click **Process Files**\n"
        "- Ask anything about your docs\n"
        "- Answers cite source files"
    )


# ── Main area ────────────────────────────────────────────────────────────────
st.markdown(
    """
    <div class="main-header">
        <h1>📚 DocuMind AI</h1>
        <p>Upload your documents and chat with them using Groq-powered intelligence</p>
    </div>
    """,
    unsafe_allow_html=True,
)

if st.session_state.loaded_files:
    cols = st.columns(3)
    with cols[0]:
        st.markdown(
            f'<div class="stat-card"><div class="value">{len(st.session_state.loaded_files)}</div>'
            f'<div class="label">Documents</div></div>',
            unsafe_allow_html=True,
        )
    with cols[1]:
        st.markdown(
            f'<div class="stat-card"><div class="value">{len(st.session_state.chunks)}</div>'
            f'<div class="label">Text Chunks</div></div>',
            unsafe_allow_html=True,
        )
    with cols[2]:
        st.markdown(
            f'<div class="stat-card"><div class="value">{len(st.session_state.messages)}</div>'
            f'<div class="label">Messages</div></div>',
            unsafe_allow_html=True,
        )
    st.markdown("<br>", unsafe_allow_html=True)

chat_container = st.container()

with chat_container:
    if not st.session_state.messages:
        st.markdown(
            """
            <div style="text-align:center; padding: 3rem 1rem; color: #9ca3af;">
                <div style="font-size: 3rem;">💬</div>
                <h3 style="color:#6b7280; margin-top:0.5rem;">Start a conversation</h3>
                <p>Upload documents in the sidebar, then ask questions like:<br>
                <em>"Summarize the key points"</em> · <em>"What does section 3 say?"</em> · <em>"List all dates mentioned"</em></p>
            </div>
            """,
            unsafe_allow_html=True,
        )
    else:
        for msg in st.session_state.messages:
            avatar = "🧑" if msg["role"] == "user" else "🤖"
            with st.chat_message(msg["role"], avatar=avatar):
                st.markdown(msg["content"])

prompt = st.chat_input(
    "Ask a question about your documents...",
    disabled=not st.session_state.chunks,
)

if prompt:
    if not st.session_state.chunks:
        st.warning("Please upload and process documents first.")
    else:
        st.session_state.messages.append({"role": "user", "content": prompt})

        relevant = retrieve_relevant_chunks(prompt, st.session_state.chunks)
        context = build_context(relevant)

        system_prompt = (
            "You are DocuMind, a helpful document assistant. "
            "Answer questions using ONLY the provided document context. "
            "If the answer is not in the context, say you could not find it in the uploaded documents. "
            "Be clear, concise, and mention which source file(s) your answer comes from when possible.\n\n"
            f"DOCUMENT CONTEXT:\n{context}"
        )

        api_messages = [{"role": "system", "content": system_prompt}]
        for msg in st.session_state.messages[:-1]:
            api_messages.append({"role": msg["role"], "content": msg["content"]})
        api_messages.append({"role": "user", "content": prompt})

        with st.chat_message("assistant", avatar="🤖"):
            with st.spinner("Thinking..."):
                try:
                    answer = ask_groq(api_messages)
                except Exception as exc:
                    answer = f"Sorry, something went wrong: {exc}"

            st.markdown(answer)

        st.session_state.messages.append({"role": "assistant", "content": answer})
        st.rerun()
